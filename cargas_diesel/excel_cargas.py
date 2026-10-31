"""
Excel del listado de cargas (Next: /diesel/cargas, botón "Exportar Excel").

Una tabla de Excel (con filtros y renglones alternados), todo centrado, con
hipervínculos "Ver foto" a cada foto de la carga y una fila de totales.
Usa los mismos filtros que el listado. El rendimiento solo aplica a 3R
RECYCLING y 3R TRANSPORTE: los permisionarios son unidades ajenas.
La pestaña "Por permisionario" trae los totales, una tabla por permisionario y las
compras y ajustes de diésel del periodo.
"""
import io

from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import AjusteInventario, CompraCombustible, maneja_rendimiento
from .views import SIN_PERMISIONARIO, filtrar_cargas, get_patios_usuario

MAX_FILAS = 20000
# Sin cinchos ni fotos de horas del thermo, sticker y motor: ya no se capturan.
# Con los mismos nombres que en la app (captura y bitácora de cargas).
FOTOS = [('foto_bomba', 'Foto inicio (Bomba)'), ('foto_bomba_final', 'Foto final (Bomba)'),
         ('foto_thermo', 'Foto inicio (Thermo)'), ('foto_thermo_final', 'Foto final (Thermo)'), ('foto_odometro', 'Foto odómetro')]
COMPLEMENTARIAS = 3  # columnas para las fotos complementarias (las demás se ven en la app)


def _url(request, archivo):
    if not archivo:
        return None
    try:
        url = archivo.url
    except Exception:
        return None
    return url if url.startswith('http') else request.build_absolute_uri(url)


def compras_y_ajustes(user, params):
    """Compras y ajustes de diésel de los patios del usuario con el patio y las fechas del listado
    (no dependen del permisionario ni de la unidad)."""
    filtros = {'patio__in': get_patios_usuario(user)}
    if params.get('patio'):
        filtros['patio_id'] = params['patio']
    if params.get('fecha_inicio'):
        filtros['fecha__date__gte'] = params['fecha_inicio']
    if params.get('fecha_fin'):
        filtros['fecha__date__lte'] = params['fecha_fin']
    compras = CompraCombustible.objects.filter(tipo='diesel', **filtros).select_related('patio', 'usuario').order_by('fecha', 'pk')
    ajustes = AjusteInventario.objects.filter(tipo_combustible='diesel', **filtros).select_related('patio', 'usuario').order_by('fecha', 'pk')
    return list(compras[:MAX_FILAS]), list(ajustes[:MAX_FILAS])


def _hoja_permisionarios(wb, cargas, subtitulo, compras, ajustes):
    """Pestaña "Por permisionario": los totales de cada permisionario, una tabla por permisionario
    con sus unidades y thermos, y las compras y ajustes de diésel del periodo (tablas con filtros)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.table import Table, TableStyleInfo

    perms, equipos = {}, {}
    for c in cargas:
        nombre = c.permisionario_nombre or 'Sin permisionario'
        thermo = c.tipo_carga == 'thermo'
        litros = float(c.litros_diesel + c.litros_thermo)
        p = perms.setdefault(nombre, {'unidades': set(), 'thermos': set(), 'cargas': 0, 'litros_u': 0.0, 'litros_t': 0.0, 'costo': 0.0})
        p['thermos' if thermo else 'unidades'].add(c.equipo)
        p['cargas'] += 1
        p['litros_u'] += float(c.litros_diesel)
        p['litros_t'] += float(c.litros_thermo)
        p['costo'] += float(c.costo_total)
        e = equipos.setdefault((nombre, 'Thermo' if thermo else 'Unidad', c.equipo), {'cargas': 0, 'litros': 0.0, 'costo': 0.0, 'ultima': None})
        e['cargas'] += 1
        e['litros'] += litros
        e['costo'] += float(c.costo_total)
        e['ultima'] = max(e['ultima'] or c.fecha_carga, c.fecha_carga)
    orden = sorted(perms, key=lambda n: (-(perms[n]['litros_u'] + perms[n]['litros_t']), n))

    ws = wb.create_sheet('Por permisionario')
    ws.sheet_view.showGridLines = False
    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    LITROS, DINERO = '#,##0.00', '"$"#,##0.00'

    def encabezado(fila, columnas):
        for i, (titulo, ancho, _) in enumerate(columnas, 1):
            celda = ws.cell(row=fila, column=i, value=titulo)
            celda.font, celda.fill, celda.alignment = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1D4ED8'), centro
            ws.column_dimensions[L(i)].width = max(ws.column_dimensions[L(i)].width or 0, ancho)
        ws.row_dimensions[fila].height = 30

    def filas(inicio, columnas, valores):
        for r, fila in enumerate(valores, inicio):
            for i, v in enumerate(fila, 1):
                celda = ws.cell(row=r, column=i, value=v)
                celda.alignment = centro
                if columnas[i - 1][2] and v is not None:
                    celda.number_format = columnas[i - 1][2]
        return inicio + len(valores) - 1

    def tabla(nombre, enc, ultima, columnas, sumar):
        t = Table(displayName=nombre, ref=f'A{enc}:{L(len(columnas))}{ultima}')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(t)
        total = ultima + 1  # totales de lo que se ve (respetan los filtros de la tabla)
        ws.cell(row=total, column=1, value='Total')
        for i, (titulo, _, fmt) in enumerate(columnas, 1):
            if titulo in sumar:
                ws.cell(row=total, column=i, value=f'=SUBTOTAL(109,{L(i)}{enc + 1}:{L(i)}{ultima})').number_format = fmt
            celda = ws.cell(row=total, column=i)
            celda.font, celda.alignment = Font(bold=True), centro
            celda.fill = PatternFill('solid', fgColor='EEF2F7')
            celda.border = Border(top=Side(style='medium', color='1D4ED8'))
        return total

    def titulo(fila, texto, tamano=13, color='0B1730'):
        ws.cell(row=fila, column=1, value=texto).font = Font(bold=True, size=tamano, color=color)

    def nota(fila, texto):
        ws.cell(row=fila, column=1, value=texto).font = Font(italic=True, size=10, color='475569')

    def seccion(fila, texto, nombre_tabla, columnas, valores, sumar, vacio):
        """Título, tabla y fila de total; si no hay filas, solo el aviso. Devuelve la última fila usada."""
        titulo(fila, texto)
        if not valores:
            nota(fila + 1, vacio)
            return fila + 1
        encabezado(fila + 1, columnas)
        return tabla(nombre_tabla, fila + 1, filas(fila + 2, columnas, valores), columnas, sumar)

    titulo(1, 'Desglose por permisionario, compras y ajustes · 3R Recycling', 15)
    nota(2, subtitulo)
    fila = 4

    # 1) Totales de cada permisionario
    if orden:
        totales = [('Permisionario', 28, None), ('Unidades', 11, '0'), ('Thermos', 11, '0'), ('Cargas', 10, '0'),
                   ('Litros unidades', 16, LITROS), ('Litros thermos', 16, LITROS), ('Total litros', 16, LITROS), ('Costo', 16, DINERO)]
        fila = seccion(fila, 'Totales por permisionario', 'TotalesPermisionario', totales,
                       [[n, len(perms[n]['unidades']), len(perms[n]['thermos']), perms[n]['cargas'], perms[n]['litros_u'],
                         perms[n]['litros_t'], perms[n]['litros_u'] + perms[n]['litros_t'], perms[n]['costo']] for n in orden],
                       {'Unidades', 'Thermos', 'Cargas', 'Litros unidades', 'Litros thermos', 'Total litros', 'Costo'}, '')

        # 2) Una tabla por permisionario con sus unidades y sus thermos
        fila += 3
        titulo(fila, 'Detalle por permisionario', 14)
        detalle = [('Tipo', 11, None), ('Unidad / Thermo', 24, None), ('Cargas', 10, '0'), ('Litros', 16, LITROS),
                   ('% del permisionario', 16, '0.0%'), ('Costo', 16, DINERO), ('Última carga', 16, None)]
        for k, n in enumerate(orden, 1):
            p = perms[n]
            total_perm = p['litros_u'] + p['litros_t']
            suyos = sorted(((e, v) for e, v in equipos.items() if e[0] == n), key=lambda ev: (ev[0][1] != 'Unidad', -ev[1]['litros'], ev[0][2]))
            def cuenta(k, uno, varios):
                return f'{k} {uno if k == 1 else varios}'
            fila = seccion(
                fila + 2, f'{n} · {cuenta(len(p["unidades"]), "unidad", "unidades")} · {cuenta(len(p["thermos"]), "thermo", "thermos")}'
                          f' · {cuenta(p["cargas"], "carga", "cargas")} · {total_perm:,.2f} L',
                f'Permisionario{k}', detalle,
                [[tipo, equipo, v['cargas'], v['litros'], v['litros'] / total_perm if total_perm else None, v['costo'],
                  timezone.localtime(v['ultima']).strftime('%d/%m/%Y')] for (_, tipo, equipo), v in suyos],
                {'Cargas', 'Litros', 'Costo'}, '')

    # 3) Compras y 4) ajustes de diésel del periodo (por patio: no dependen del permisionario)
    fila += 3
    nota(fila + 1, 'Compras y ajustes de diésel de los patios y fechas del reporte (no dependen del permisionario ni de la unidad).')
    compras_cols = [('Fecha', 17, None), ('Patio', 20, None), ('Proveedor', 26, None), ('Factura', 14, None),
                    ('Litros', 14, LITROS), ('Precio/L', 13, DINERO), ('Total', 16, DINERO), ('Registró', 22, None)]
    fila = seccion(fila + 2, 'Compras de diésel', 'Compras', compras_cols,
                   [[timezone.localtime(c.fecha).strftime('%d/%m/%Y %H:%M'), c.patio.nombre, c.proveedor_nombre or '—',
                     c.numero_factura or '—', float(c.litros), float(c.precio_unitario), float(c.total),
                     c.usuario.get_full_name() or c.usuario.username] for c in compras],
                   {'Litros', 'Total'}, 'Sin compras de diésel en el periodo.')
    ajustes_cols = [('Fecha', 17, None), ('Patio', 20, None), ('Movimiento', 12, None), ('Litros (+/−)', 14, LITROS),
                    ('Motivo', 40, None), ('Antes', 13, LITROS), ('Después', 13, LITROS), ('Registró', 22, None)]
    seccion(fila + 3, 'Ajustes de inventario', 'Ajustes', ajustes_cols,
            [[timezone.localtime(a.fecha).strftime('%d/%m/%Y %H:%M'), a.patio.nombre,
              'Entrada' if a.tipo_movimiento == 'entrada' else 'Salida',
              float(a.litros) * (1 if a.tipo_movimiento == 'entrada' else -1), a.motivo or '—',
              float(a.cantidad_anterior), float(a.cantidad_posterior), a.usuario.get_full_name() or a.usuario.username]
             for a in ajustes],
            {'Litros (+/−)'}, 'Sin ajustes de diésel en el periodo.')


class CargasExcelAPIView(APIView):
    """GET con los filtros del listado (?patio, ?unidad, ?permisionario, ?fecha_inicio, ?fecha_fin): descarga .xlsx."""

    def get(self, request):
        user = request.user
        if not (user.is_superuser or user.has_perm('ternium.acceso_diesel')):
            return Response({'detail': 'No tienes acceso a Control Diésel.'}, status=status.HTTP_403_FORBIDDEN)
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter as L
        from openpyxl.worksheet.table import Table, TableStyleInfo

        cargas = list(filtrar_cargas(user, request.query_params).order_by('-fecha_carga', '-pk')[:MAX_FILAS])
        # (título, ancho, formato numérico o None)
        columnas = [
            ('#', 6, '0'), ('Fecha', 12, None), ('Hora', 8, None), ('Tipo', 9, None), ('Unidad / Thermo', 16, None),
            ('Permisionario', 26, None), ('Operador', 24, None), ('Patio', 20, None), ('Odómetro (km)', 15, '#,##0'),
            ('Litros diésel', 13, '#,##0.00'), ('Litros thermo', 13, '#,##0.00'), ('Total litros', 13, '#,##0.00'),
            ('Precio/L', 12, '"$"#,##0.00'), ('Costo', 15, '"$"#,##0.00'), ('Rendimiento (km/L)', 17, '0.00'),
            ('Despachador', 22, None),
            *[(nombre, 16, None) for _, nombre in FOTOS],
            *[(f'Foto complementaria {n}', 16, None) for n in range(1, COMPLEMENTARIAS + 1)],
        ]
        wb = Workbook()
        ws = wb.active
        ws.title = 'Cargas'
        ws.sheet_view.showGridLines = False
        centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
        fino = Side(style='thin', color='D7DEE8')

        # Título y filtros
        ultima_col = L(len(columnas))
        ws.merge_cells(f'A1:{ultima_col}1')
        ws['A1'] = 'Listado de cargas de diésel · 3R Recycling'
        ws['A1'].font = Font(bold=True, size=15, color='FFFFFF')
        ws['A1'].fill = PatternFill('solid', fgColor='0B1730')
        ws['A1'].alignment = centro
        ws.row_dimensions[1].height = 30
        p = request.query_params
        perm = (p.get('permisionario') or '').strip()
        filtros = [f'Desde {p["fecha_inicio"]}' if p.get('fecha_inicio') else '', f'hasta {p["fecha_fin"]}' if p.get('fecha_fin') else '',
                   f'Unidad/thermo: {p["unidad"]}' if p.get('unidad') else '',
                   ('Sin permisionario' if perm == SIN_PERMISIONARIO else f'Permisionario: {perm}') if perm else '',
                   {'verificado': 'Auditoría: Verificado', 'alerta': 'Auditoría: Alerta', 'sin': 'Auditoría: sin estatus'}.get(p.get('auditoria'), '')]
        ws.merge_cells(f'A2:{ultima_col}2')
        subtitulo = (' · '.join(f for f in filtros if f) or 'Todas las cargas') + \
            f' · {len(cargas)} cargas · generado por {user.get_full_name() or user.username} el {timezone.localtime():%d/%m/%Y %H:%M}'
        ws['A2'] = subtitulo
        ws['A2'].font = Font(italic=True, size=10, color='475569')
        ws['A2'].alignment = centro

        enc = 4
        for i, (titulo, ancho, _) in enumerate(columnas, 1):
            celda = ws.cell(row=enc, column=i, value=titulo)
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='1D4ED8')
            celda.alignment = centro
            ws.column_dimensions[L(i)].width = ancho
        ws.row_dimensions[enc].height = 30

        primera_foto = len(columnas) - len(FOTOS) - COMPLEMENTARIAS + 1
        for n, c in enumerate(cargas, 1):
            r = enc + n
            local = timezone.localtime(c.fecha_carga)
            rinde = maneja_rendimiento(c.permisionario_nombre)
            thermo = c.tipo_carga == 'thermo'
            valores = [
                n, local.strftime('%d/%m/%Y'), local.strftime('%H:%M'), 'Thermo' if thermo else 'Unidad', c.equipo,
                c.permisionario_nombre or 'Sin permisionario', c.operador_nombre or None, c.patio.nombre,
                float(c.odometro) if c.odometro is not None and not thermo else None,
                float(c.litros_diesel), float(c.litros_thermo), float(c.litros_diesel + c.litros_thermo),
                float(c.precio_litro) or None, float(c.costo_total),
                float(c.rendimiento) if rinde and c.rendimiento is not None else None,
                c.usuario.get_full_name() or c.usuario.username,
            ]
            for i, v in enumerate(valores, 1):
                celda = ws.cell(row=r, column=i, value=v if v is not None else '—')
                if columnas[i - 1][2] and v is not None:
                    celda.number_format = columnas[i - 1][2]
            archivos = [getattr(c, campo) for campo, _ in FOTOS]
            extras = [f.foto for f in c.fotos_complementarias.all()][:COMPLEMENTARIAS]
            archivos += extras + [None] * (COMPLEMENTARIAS - len(extras))
            for k, archivo in enumerate(archivos):
                celda = ws.cell(row=r, column=primera_foto + k)
                url = _url(request, archivo)
                if url:
                    celda.value = 'Ver foto'
                    celda.hyperlink = url
                    celda.font = Font(color='1D4ED8', underline='single', bold=True)
                else:
                    celda.value = '—'
            for i in range(1, len(columnas) + 1):
                ws.cell(row=r, column=i).alignment = centro

        ultima = enc + len(cargas)
        if cargas:
            # Tabla de Excel: filtros en el encabezado y renglones alternados
            tabla = Table(displayName='Cargas', ref=f'A{enc}:{ultima_col}{ultima}')
            tabla.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(tabla)
            t = ultima + 1
            ws.cell(row=t, column=5, value='Totales')
            for i, (titulo, _, fmt) in enumerate(columnas, 1):
                if titulo in ('Litros diésel', 'Litros thermo', 'Total litros', 'Costo'):
                    ws.cell(row=t, column=i, value=f'=SUBTOTAL(109,{L(i)}{enc + 1}:{L(i)}{ultima})').number_format = fmt
                celda = ws.cell(row=t, column=i)
                celda.font = Font(bold=True)
                celda.alignment = centro
                celda.fill = PatternFill('solid', fgColor='EEF2F7')
                celda.border = Border(top=Side(style='medium', color='1D4ED8'), bottom=fino)
            ws.row_dimensions[t].height = 22
        else:
            ws.cell(row=enc + 1, column=1, value='No hay cargas con esos filtros.')
        ws.freeze_panes = ws.cell(row=enc + 1, column=6)
        compras, ajustes = compras_y_ajustes(user, request.query_params)
        if cargas or compras or ajustes:
            _hoja_permisionarios(wb, cargas, subtitulo, compras, ajustes)

        salida = io.BytesIO()
        wb.save(salida)
        nombre = f'Cargas de diesel {timezone.localdate():%Y-%m-%d}.xlsx'
        respuesta = HttpResponse(salida.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
        return respuesta
