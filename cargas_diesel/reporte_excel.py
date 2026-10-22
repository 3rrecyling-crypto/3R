"""
Reporte en Excel de diésel a la medida de lo que pide el usuario (lo arma la IA
del tablero con la herramienta `generar_reporte_excel`; se descarga desde el chat).

La especificación del reporte (periodo, patio, secciones y filtros) viene de la
IA y se vuelve a validar aquí al descargar; los datos siempre salen de los
patios del usuario. El costo de cada carga es litros × precio por litro del
tótem del que salió; los permisionarios no manejan rendimiento.
"""
import io
import re
from collections import defaultdict
from datetime import timedelta

from django.db.models import Q
from django.utils import timezone

from .models import AjusteInventario, CargaDiesel, CompraCombustible, es_3r_nombre
from .views import get_patios_usuario
from .views_tablero import rango

SECCIONES = {
    'por_dia': 'Litros por día',
    'equipos_3r': 'Equipos 3R',
    'cargas_3r': 'Cargas 3R',
    'permisionarios': 'Permisionarios',
    'cargas_permisionarios': 'Cargas permisionarios',
    'compras': 'Compras',
    'ajustes': 'Ajustes',
    'totems': 'Tótems',
}
MAX_FILAS = 5000
AZUL, AZUL_OSCURO, GRIS = '1D4ED8', '0B1730', 'EEF2F7'


def normalizar(entrada):
    """Especificación limpia del reporte (de la IA o del navegador)."""
    e = entrada if isinstance(entrada, dict) else {}
    desde, hasta = rango(str(e.get('desde') or ''), str(e.get('hasta') or ''))
    secciones = [s for s in (e.get('secciones') or []) if s in SECCIONES] or list(SECCIONES)
    patio = e.get('patio_id')
    return {
        'titulo': ' '.join(str(e.get('titulo') or 'Reporte de diésel').split())[:80],
        'desde': desde.isoformat(), 'hasta': hasta.isoformat(),
        'patio_id': int(patio) if str(patio or '').isdigit() and int(patio) > 0 else None,
        'secciones': list(dict.fromkeys(secciones)),
        'permisionario': ' '.join(str(e.get('permisionario') or '').split())[:100],
        'equipo': ' '.join(str(e.get('equipo') or '').split())[:50],
    }


def _hora(f):
    return timezone.localtime(f).strftime('%Y-%m-%d %H:%M')


def datos(user, spec):
    """Todo lo que lleva el reporte, con los filtros de la especificación."""
    desde, hasta = rango(spec['desde'], spec['hasta'])
    patios = get_patios_usuario(user)
    if spec['patio_id']:
        patios = patios.filter(pk=spec['patio_id'])
    qs = CargaDiesel.objects.filter(patio__in=patios, fecha_carga__date__gte=desde, fecha_carga__date__lte=hasta) \
        .select_related('unidad', 'thermo', 'patio__totem', 'usuario').order_by('fecha_carga', 'pk')
    if spec['permisionario']:
        qs = qs.filter(permisionario_nombre__icontains=spec['permisionario'])
    if spec['equipo']:
        qs = qs.filter(Q(unidad__numero_economico__iexact=spec['equipo']) | Q(thermo__identificador__iexact=spec['equipo']))
    cargas = list(qs[:MAX_FILAS])

    dias = [desde + timedelta(days=n) for n in range((hasta - desde).days + 1)]
    por_dia = {d: [0.0, 0.0, 0.0] for d in dias}  # 3R unidades, 3R thermos, permisionarios
    filas_3r, filas_perm = [], []
    equipos = defaultdict(lambda: {'cargas': 0, 'litros': 0.0, 'km': 0.0, 'litros_rend': 0.0, 'costo': 0.0, 'ultima': ''})
    perms = defaultdict(lambda: {'cargas': 0, 'unidades': set(), 'thermos': set(), 'litros': 0.0, 'costo': 0.0})
    for c in cargas:
        litros = float(c.litros_diesel + c.litros_thermo)
        es3r = es_3r_nombre(c.permisionario_nombre)
        dia = por_dia.get(timezone.localtime(c.fecha_carga).date())
        fila = {
            'fecha': _hora(c.fecha_carga), 'equipo': c.equipo, 'tipo': 'Thermo' if c.tipo_carga == 'thermo' else 'Unidad',
            'permisionario': c.permisionario_nombre or 'Sin permisionario', 'operador': c.operador_nombre or '—',
            'patio': c.patio.nombre,
            'lectura': float(c.hrs_thermo) if c.tipo_carga == 'thermo' else (float(c.odometro) if c.odometro is not None else None),
            'litros': litros, 'precio': float(c.precio_litro) or None, 'costo': float(c.costo_total),
            'rendimiento': float(c.rendimiento) if es3r and c.rendimiento is not None else None,
            'despachador': c.usuario.get_full_name() or c.usuario.username,
        }
        if es3r:
            filas_3r.append(fila)
            if dia:
                dia[0] += float(c.litros_diesel)
                dia[1] += float(c.litros_thermo)
            e = equipos[(fila['tipo'], c.equipo)]
            e['cargas'] += 1
            e['litros'] += litros
            e['costo'] += fila['costo']
            e['ultima'] = fila['fecha']
            if c.rendimiento is not None and c.litros_diesel:
                e['km'] += float(c.rendimiento) * float(c.litros_diesel)
                e['litros_rend'] += float(c.litros_diesel)
        else:
            filas_perm.append(fila)
            if dia:
                dia[2] += litros
            p = perms[fila['permisionario']]
            p['cargas'] += 1
            p['litros'] += litros
            p['costo'] += fila['costo']
            (p['thermos'] if c.tipo_carga == 'thermo' else p['unidades']).add(c.equipo)

    total_perm = sum(p['litros'] for p in perms.values()) or 1
    rango_fechas = {'patio__in': patios, 'fecha__date__gte': desde, 'fecha__date__lte': hasta}
    compras = list(CompraCombustible.objects.filter(tipo='diesel', **rango_fechas).select_related('patio', 'usuario').order_by('fecha')[:MAX_FILAS])
    ajustes = list(AjusteInventario.objects.filter(tipo_combustible='diesel', **rango_fechas).select_related('patio', 'usuario').order_by('fecha')[:MAX_FILAS])
    totems = []
    for p in patios.select_related('totem'):
        t = getattr(p, 'totem', None)
        if t:
            ultima = CompraCombustible.objects.filter(patio=p, tipo='diesel').order_by('-fecha', '-pk').first()
            totems.append({'patio': p.nombre, 'totem': t.nombre, 'capacidad': float(t.capacidad_diesel),
                           'cantidad': float(t.cantidad_diesel), 'porcentaje': t.porcentaje_diesel(),
                           'precio': float(ultima.precio_unitario) if ultima else None})
    return {
        'desde': desde, 'hasta': hasta, 'patio': patios.first().nombre if spec['patio_id'] and patios.exists() else 'Todos los patios',
        'truncado': len(cargas) >= MAX_FILAS,
        'cargas_3r': filas_3r, 'cargas_permisionarios': filas_perm,
        'por_dia': [(d, *v) for d, v in por_dia.items()],
        'equipos_3r': sorted(({
            'equipo': eq, 'tipo': tipo, 'cargas': e['cargas'], 'litros': e['litros'], 'km': e['km'],
            'rendimiento': round(e['km'] / e['litros_rend'], 2) if e['litros_rend'] else None,
            'costo': e['costo'], 'precio': e['costo'] / e['litros'] if e['litros'] and e['costo'] else None, 'ultima': e['ultima'],
        } for (tipo, eq), e in equipos.items()), key=lambda x: -x['litros']),
        'permisionarios': sorted(({
            'nombre': n, 'cargas': p['cargas'], 'unidades': len(p['unidades']), 'thermos': len(p['thermos']),
            'litros': p['litros'], 'porcentaje': p['litros'] / total_perm, 'costo': p['costo'],
            'precio': p['costo'] / p['litros'] if p['litros'] and p['costo'] else None,
        } for n, p in perms.items()), key=lambda x: -x['litros']),
        'compras': [{'fecha': _hora(c.fecha), 'patio': c.patio.nombre, 'proveedor': c.proveedor_nombre or '—', 'litros': float(c.litros),
                     'precio': float(c.precio_unitario), 'total': float(c.total), 'factura': c.numero_factura,
                     'usuario': c.usuario.get_full_name() or c.usuario.username} for c in compras],
        'ajustes': [{'fecha': _hora(a.fecha), 'patio': a.patio.nombre, 'movimiento': 'Entrada' if a.tipo_movimiento == 'entrada' else 'Salida',
                     'litros': float(a.litros) * (1 if a.tipo_movimiento == 'entrada' else -1), 'motivo': a.motivo,
                     'antes': float(a.cantidad_anterior), 'despues': float(a.cantidad_posterior),
                     'usuario': a.usuario.get_full_name() or a.usuario.username} for a in ajustes],
        'totems': totems,
    }


def resumen(user, spec):
    """Lo que la IA necesita saber del reporte para describirlo (sin armar el archivo)."""
    d = datos(user, spec)
    todas = d['cargas_3r'] + d['cargas_permisionarios']
    return {
        'reporte': spec, 'patio': d['patio'], 'cargas': len(todas),
        'litros': round(sum(c['litros'] for c in todas), 2), 'costo': round(sum(c['costo'] for c in todas), 2),
        'cargas_3r': len(d['cargas_3r']), 'cargas_permisionarios': len(d['cargas_permisionarios']),
        'compras': len(d['compras']), 'ajustes': len(d['ajustes']),
        'hojas': ['Resumen'] + [SECCIONES[s] for s in spec['secciones']],
    }


def nombre_archivo(spec):
    base = re.sub(r'[^\w\s-]', '', spec['titulo']).strip()[:60] or 'Reporte de diesel'
    return f'{base} {spec["desde"]} a {spec["hasta"]}.xlsx'


def construir(user, spec, pregunta='', respuesta='') -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    d = datos(user, spec)
    wb = Workbook()
    fino = Side(style='thin', color='D7DEE8')
    FMT = {'L': '#,##0" L"', '$': '"$"#,##0.00', '$L': '"$"#,##0.00"/L"', 'n': '#,##0', 'km': '#,##0" km"',
           'r': '0.00" km/L"', '%': '0.0%', 'h': '#,##0.0" h"', 't': '@'}

    def encabezado(ws, titulo, ancho):
        ws.sheet_view.showGridLines = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ancho, 4))
        ws['A1'] = titulo
        ws['A1'].font = Font(bold=True, size=15, color='FFFFFF')
        ws['A1'].fill = PatternFill('solid', fgColor=AZUL_OSCURO)
        ws['A1'].alignment = Alignment(vertical='center', indent=1)
        ws.row_dimensions[1].height = 30
        ws['A2'] = (f'3R Recycling · Control Diésel · {d["desde"]:%d/%m/%Y} – {d["hasta"]:%d/%m/%Y} · {d["patio"]}'
                    + (f' · Permisionario: {spec["permisionario"]}' if spec['permisionario'] else '')
                    + (f' · Equipo: {spec["equipo"]}' if spec['equipo'] else ''))
        ws['A2'].font = Font(italic=True, size=10, color='475569')
        return 4

    def tabla(ws, fila, columnas, renglones, totales=()):
        """columnas: [(titulo, clave, formato, ancho)]. Encabezado azul, filtros, fila de totales con SUMA."""
        for i, (titulo, _, _, ancho) in enumerate(columnas, 1):
            celda = ws.cell(row=fila, column=i, value=titulo)
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor=AZUL)
            celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            celda.border = Border(bottom=fino)
            ws.column_dimensions[L(i)].width = max(ws.column_dimensions[L(i)].width or 0, ancho)
        ws.row_dimensions[fila].height = 22
        for r, reng in enumerate(renglones, fila + 1):
            for i, (_, clave, fmt, _) in enumerate(columnas, 1):
                celda = ws.cell(row=r, column=i, value=reng.get(clave))
                celda.number_format = FMT.get(fmt, 'General')
                celda.border = Border(bottom=fino)
                if r % 2 == 0:
                    celda.fill = PatternFill('solid', fgColor='F8FAFC')
        ultima = fila + len(renglones)
        if renglones:
            ws.auto_filter.ref = f'A{fila}:{L(len(columnas))}{ultima}'
            ws.freeze_panes = ws.cell(row=fila + 1, column=1)
        if totales and renglones:
            t = ultima + 1
            ws.cell(row=t, column=1, value='Total').font = Font(bold=True)
            for i, (_, clave, fmt, _) in enumerate(columnas, 1):
                if clave in totales:
                    celda = ws.cell(row=t, column=i, value=f'=SUM({L(i)}{fila + 1}:{L(i)}{ultima})')
                    celda.number_format = FMT.get(fmt, 'General')
                    celda.font = Font(bold=True)
                ws.cell(row=t, column=i).fill = PatternFill('solid', fgColor=GRIS)
                ws.cell(row=t, column=i).border = Border(top=Side(style='medium', color=AZUL))
            ultima = t
        return ultima

    # ── Resumen ──
    ws = wb.active
    ws.title = 'Resumen'
    fila = encabezado(ws, spec['titulo'], 6)
    todas = d['cargas_3r'] + d['cargas_permisionarios']
    litros = sum(c['litros'] for c in todas)
    costo = sum(c['costo'] for c in todas)
    con_precio = sum(c['litros'] for c in todas if c['precio'])
    km = sum(e['km'] for e in d['equipos_3r'])
    litros_rend = sum(e['km'] / e['rendimiento'] for e in d['equipos_3r'] if e['rendimiento'])
    indicadores = [
        ('Litros despachados', litros, 'L'), ('Cargas', len(todas), 'n'), ('Costo (precio del tótem)', costo, '$'),
        ('Precio promedio por litro', costo / con_precio if con_precio else None, '$L'),
        ('Litros 3R RECYCLING', sum(c['litros'] for c in d['cargas_3r']), 'L'),
        ('Rendimiento 3R (ponderado)', km / litros_rend if litros_rend else None, 'r'),
        ('Km recorridos 3R', km, 'km'),
        ('Litros permisionarios', sum(c['litros'] for c in d['cargas_permisionarios']), 'L'),
        ('Compras de diésel', sum(c['total'] for c in d['compras']), '$'),
        ('Litros comprados', sum(c['litros'] for c in d['compras']), 'L'),
        ('Ajustes netos', sum(a['litros'] for a in d['ajustes']), 'L'),
    ]
    fila = tabla(ws, fila, [('Indicador', 'k', 't', 34), ('Valor', 'v', 'n', 22)], [])
    for i, (k, v, fmt) in enumerate(indicadores, fila + 1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color='334155')
        celda = ws.cell(row=i, column=2, value=v if v is not None else '—')
        celda.number_format = FMT[fmt]
        celda.font = Font(bold=True, size=12)
        for col in (1, 2):
            ws.cell(row=i, column=col).border = Border(bottom=fino)
    fila += len(indicadores) + 2
    if respuesta:
        ws.cell(row=fila, column=1, value='Lo que pidió' if pregunta else 'Análisis').font = Font(bold=True, color=AZUL)
        if pregunta:
            ws.cell(row=fila + 1, column=1, value=pregunta).font = Font(italic=True)
            fila += 2
            ws.cell(row=fila, column=1, value='Análisis de la IA').font = Font(bold=True, color=AZUL)
        texto = respuesta.replace('**', '')
        ws.merge_cells(start_row=fila + 1, start_column=1, end_row=fila + 1, end_column=6)
        celda = ws.cell(row=fila + 1, column=1, value=texto[:32000])
        celda.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[fila + 1].height = min(409, 15 * (texto.count('\n') + 2 + len(texto) // 110))
        fila += 3
    ws.cell(row=fila, column=1, value=f'Generado por {user.get_full_name() or user.username} el {timezone.localtime():%d/%m/%Y %H:%M}'
            + (' · se limitó a las primeras 5,000 cargas' if d['truncado'] else '')).font = Font(size=9, color='94A3B8')

    # ── Secciones pedidas ──
    hojas = {}
    for s in spec['secciones']:
        hojas[s] = h = wb.create_sheet(SECCIONES[s])
        f = encabezado(h, f'{spec["titulo"]} · {SECCIONES[s]}', 8)
        if s == 'por_dia':
            ult = tabla(h, f, [('Fecha', 'fecha', 't', 14), ('3R unidades', 'u', 'L', 16), ('3R thermos', 't', 'L', 16),
                               ('Permisionarios', 'p', 'L', 16), ('Total', 'tot', 'L', 16)],
                        [{'fecha': dia.strftime('%d/%m/%Y'), 'u': u, 't': t, 'p': p, 'tot': u + t + p} for dia, u, t, p in d['por_dia']],
                        totales=('u', 't', 'p', 'tot'))
            if d['por_dia']:
                g = BarChart()
                g.type, g.grouping, g.overlap = 'col', 'stacked', 100
                g.title, g.y_axis.title = 'Litros por día', 'Litros'
                g.height, g.width = 9, 24
                g.add_data(Reference(h, min_col=2, max_col=4, min_row=f, max_row=ult - 1), titles_from_data=True)
                g.set_categories(Reference(h, min_col=1, min_row=f + 1, max_row=ult - 1))
                for serie, color in zip(g.series, ('2A78D6', 'EB6834', '1BAF7A')):
                    serie.graphicalProperties.solidFill = color
                    serie.graphicalProperties.line.solidFill = color
                h.add_chart(g, f'G{f}')
        elif s == 'equipos_3r':
            tabla(h, f, [('Equipo', 'equipo', 't', 14), ('Tipo', 'tipo', 't', 10), ('Cargas', 'cargas', 'n', 10),
                         ('Litros', 'litros', 'L', 14), ('Km recorridos', 'km', 'km', 16), ('Rendimiento', 'rendimiento', 'r', 15),
                         ('Precio prom.', 'precio', '$L', 15), ('Costo', 'costo', '$', 16), ('Última carga', 'ultima', 't', 18)],
                  d['equipos_3r'], totales=('cargas', 'litros', 'km', 'costo'))
        elif s in ('cargas_3r', 'cargas_permisionarios'):
            columnas = [('Fecha', 'fecha', 't', 17), ('Equipo', 'equipo', 't', 12), ('Tipo', 'tipo', 't', 9),
                        ('Permisionario', 'permisionario', 't', 24), ('Operador', 'operador', 't', 22),
                        ('Patio / tótem', 'patio', 't', 20),
                        ('Km / horas', 'lectura', 'n', 13), ('Litros', 'litros', 'L', 12),
                        *([('Rendimiento', 'rendimiento', 'r', 14)] if s == 'cargas_3r' else []),
                        ('Precio/L tótem', 'precio', '$L', 15), ('Costo', 'costo', '$', 15),
                        ('Despachador', 'despachador', 't', 20)]
            tabla(h, f, columnas, d[s], totales=('litros', 'costo'))
        elif s == 'permisionarios':
            tabla(h, f, [('Permisionario', 'nombre', 't', 28), ('Cargas', 'cargas', 'n', 10), ('Unidades', 'unidades', 'n', 10),
                         ('Thermos', 'thermos', 'n', 10), ('Litros', 'litros', 'L', 14), ('% del total', 'porcentaje', '%', 12),
                         ('Precio prom.', 'precio', '$L', 15), ('Costo', 'costo', '$', 16)],
                  d['permisionarios'], totales=('cargas', 'litros', 'costo'))
        elif s == 'compras':
            tabla(h, f, [('Fecha', 'fecha', 't', 17), ('Patio / tótem', 'patio', 't', 20), ('Proveedor', 'proveedor', 't', 24),
                         ('Litros', 'litros', 'L', 14), ('Precio/L', 'precio', '$L', 14), ('Total', 'total', '$', 16),
                         ('Factura', 'factura', 't', 14), ('Registró', 'usuario', 't', 20)],
                  d['compras'], totales=('litros', 'total'))
        elif s == 'ajustes':
            tabla(h, f, [('Fecha', 'fecha', 't', 17), ('Patio / tótem', 'patio', 't', 20), ('Movimiento', 'movimiento', 't', 12),
                         ('Litros (+/−)', 'litros', 'L', 14), ('Motivo', 'motivo', 't', 40), ('Antes', 'antes', 'L', 13),
                         ('Después', 'despues', 'L', 13), ('Registró', 'usuario', 't', 20)],
                  d['ajustes'], totales=('litros',))
        elif s == 'totems':
            tabla(h, f, [('Patio', 'patio', 't', 22), ('Tótem', 'totem', 't', 20), ('Capacidad', 'capacidad', 'L', 14),
                         ('Nivel actual', 'cantidad', 'L', 14), ('% lleno', 'porcentaje', 'n', 10), ('Precio vigente', 'precio', '$L', 15)],
                  d['totems'])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
