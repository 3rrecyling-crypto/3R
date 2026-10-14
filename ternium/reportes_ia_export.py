"""
Descargas de Reportes con IA: Excel con fórmulas y presentación de PowerPoint.

Los datos siempre se vuelven a consultar en el servidor (reportes_ia.ejecutar),
con los permisos del usuario: del navegador solo llegan las consultas.

Excel:
  - Hoja "Resumen" (indicadores; toneladas y merma con fórmula) y una hoja por consulta.
  - Columnas derivadas con fórmula: toneladas = kg / 1000, merma %, diferencia,
    promedios por remisión y % de participación.
  - Fila de totales con SUBTOTAL (se recalcula al filtrar), merma alta en rojo
    (formato condicional) y gráfica en las tablas agrupadas.

Presentación:
  - Portada, respuesta con indicadores y una diapositiva por consulta (gráfica
    editable de PowerPoint o tabla con los principales renglones).
"""
from __future__ import annotations

import copy
import re
from io import BytesIO

from django.utils import timezone

from . import reportes_ia as R

AZUL_OSCURO = '2C3E50'
AZUL = '2563EB'
ROJO = 'C0392B'
GRIS = 'EEF1F4'
UMBRAL_MERMA = 5  # % a partir del cual la merma se marca en rojo (igual que en pantalla)

# Columnas que no se suman en la fila de totales
NO_ADITIVAS = {'merma_pct', 'promedio_carga_kg', 'promedio_descarga_kg', 'tiempo_carga_h',
               'tiempo_descarga_h', 'tiempo_transito_h', 'dias_antiguedad', 'precio_kg'}
PROMEDIADAS = {'tiempo_carga_h', 'tiempo_descarga_h', 'tiempo_transito_h', 'dias_antiguedad', 'precio_kg'}
PRINCIPALES = ('descarga_ton', 'descarga_kg', 'carga_ton', 'carga_kg', 'importe', 'remisiones')
FORMATOS = {'kg': '#,##0.00', 'ton': '#,##0.000', 'pct': '0.00"%"', 'entero': '#,##0',
            'fecha': 'dd/mm/yyyy', 'fechahora': 'dd/mm/yyyy hh:mm', 'hora': 'hh:mm',
            'horas': '#,##0.00', 'dinero': '"$"#,##0.00'}


def con_columnas_para_formulas(consulta):
    """Copia de la consulta con las columnas que necesitan las fórmulas (merma y
    diferencia necesitan carga, descarga y rechazo; los promedios, remisiones)."""
    c = copy.deepcopy(consulta)
    clave = 'columnas' if c['tipo'] == 'lista' else 'metricas'
    lista = c[clave]
    if any(k in lista for k in ('merma_pct', 'diferencia_kg')):
        pos = min(lista.index(k) for k in ('merma_pct', 'diferencia_kg') if k in lista)
        for k in ('carga_kg', 'descarga_kg', 'rechazo_kg'):
            if k not in lista:
                lista.insert(pos, k)
                pos += 1
    if c['tipo'] == 'agrupado' and 'remisiones' not in lista:
        lista.insert(0, 'remisiones')
    return c


def ejecutar_consultas(consultas, user, para_excel):
    salida = []
    for consulta in consultas:
        c = R.normalizar(consulta)
        if c['tipo'] == 'conversacion':
            continue
        if para_excel:
            c = con_columnas_para_formulas(c)
        salida.append((c, R.ejecutar(c, user, para_excel=para_excel)))
    return salida


def _nombre_hoja(texto, usados):
    base = re.sub(r'[\[\]:*?/\\]', ' ', texto).strip()[:28] or 'Datos'
    nombre, n = base, 2
    while nombre.lower() in usados:
        nombre = f'{base[:26]} {n}'
        n += 1
    usados.add(nombre.lower())
    return nombre


def _texto_plano(texto):
    return re.sub(r'\*\*', '', texto or '')


# ════════════════════════════════════════════════════════════════════════════
# EXCEL
# ════════════════════════════════════════════════════════════════════════════

def construir_excel(hojas, user, titulo='', pregunta='', respuesta='') -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    resumen = wb.active
    resumen.title = 'Resumen'
    usados = {'resumen'}
    destinos = []
    for i, (c, res) in enumerate(hojas, start=1):
        nombre = _nombre_hoja(res['titulo'] if len(hojas) == 1 else f'{i} {res["titulo"]}', usados)
        ws = wb.create_sheet(nombre)
        _hoja_datos(ws, c, res, user)
        destinos.append(nombre)
    _hoja_resumen(resumen, hojas, destinos, user, titulo, pregunta, respuesta)
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _encabezado_hoja(ws, titulo, user, filtros, nota=''):
    from openpyxl.styles import Font
    ws['A1'] = titulo
    ws['A1'].font = Font(bold=True, size=14, color=AZUL_OSCURO)
    nombre = user.get_full_name() or user.username
    ws['A2'] = f'Generado con Reportes con IA por {nombre} el {timezone.localtime().strftime("%d/%m/%Y %H:%M")}'
    ws['A2'].font = Font(italic=True, size=9, color='666666')
    ws['A3'] = ' · '.join(filtros)
    ws['A3'].font = Font(size=9, color='444444')
    if nota:
        ws['A4'] = nota
        ws['A4'].font = Font(size=9, bold=True, color=ROJO)


def _estilo_encabezado(celda):
    from openpyxl.styles import Alignment, Font, PatternFill
    celda.font = Font(bold=True, color='FFFFFF')
    celda.fill = PatternFill('solid', start_color=AZUL_OSCURO, end_color=AZUL_OSCURO)
    celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _valor_excel(valor, tipo):
    if tipo == 'estatus':
        return R.STATUS.get(valor, valor)
    if tipo == 'sino':
        return '' if valor is None else ('Sí' if valor else 'No')
    if tipo == 'fechahora' and valor:
        return timezone.localtime(valor).replace(tzinfo=None)
    return valor


def _hoja_datos(ws, c, res, user):
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    tabla = res.get('tabla')
    nota = ''
    if tabla and tabla['mostradas'] < tabla['total_filas']:
        nota = (f'Se incluyen {tabla["mostradas"]:,} de {tabla["total_filas"]:,} renglones; '
                'los indicadores de la hoja Resumen cuentan todos.')
    _encabezado_hoja(ws, res['titulo'], user, res['filtros_texto'], nota)
    borde = Border(bottom=Side(style='thin', color='D0D7DE'))
    gris = PatternFill('solid', start_color=GRIS, end_color=GRIS)
    enc = 6

    if not tabla:
        _hoja_totales(ws, res, enc)
        return

    columnas = list(tabla['columnas'])
    agrupado = c['tipo'] == 'agrupado'
    principal = next((m for m in PRINCIPALES if any(col['clave'] == m for col in columnas)), None) if agrupado else None
    if principal:
        columnas.append({'clave': '_participacion', 'titulo': 'Participación', 'tipo': 'participacion'})
    col = {cc['clave']: i for i, cc in enumerate(columnas, start=1)}
    filas = tabla['filas']
    primera, ultima = enc + 1, enc + max(len(filas), 1)
    total = ultima + 1

    def ref(clave, fila, absoluta=False):
        letra = L(col[clave])
        return f'${letra}${fila}' if absoluta else f'{letra}{fila}'

    def formula(clave, fila):
        """Fórmula de una columna derivada, o None si no se puede armar."""
        if clave.endswith('_ton') and clave.replace('_ton', '_kg') in col:
            return f'={ref(clave.replace("_ton", "_kg"), fila)}/1000'
        tiene = all(k in col for k in ('carga_kg', 'descarga_kg', 'rechazo_kg'))
        if clave == 'merma_pct' and tiene:
            cg, ds, rj = (ref(k, fila) for k in ('carga_kg', 'descarga_kg', 'rechazo_kg'))
            return f'=IF({cg}>0,({cg}-{ds}-{rj})/{cg}*100,0)'
        if clave == 'diferencia_kg' and tiene:
            cg, ds, rj = (ref(k, fila) for k in ('carga_kg', 'descarga_kg', 'rechazo_kg'))
            return f'={ds}+{rj}-{cg}'
        base = {'promedio_carga_kg': 'carga_kg', 'promedio_descarga_kg': 'descarga_kg'}.get(clave)
        if base and base in col and 'remisiones' in col:
            return f'=IF({ref("remisiones", fila)}>0,{ref(base, fila)}/{ref("remisiones", fila)},0)'
        if clave == '_participacion':
            return f'=IFERROR({ref(principal, fila)}/{ref(principal, total, True)},0)'
        return None

    for i, cc in enumerate(columnas, start=1):
        _estilo_encabezado(ws.cell(row=enc, column=i, value=cc['titulo']))
    ws.row_dimensions[enc].height = 30

    for n, datos in enumerate(filas):
        fila = primera + n
        for cc in columnas:
            celda = ws.cell(row=fila, column=col[cc['clave']])
            celda.value = formula(cc['clave'], fila) or _valor_excel(datos.get(cc['clave']), cc['tipo'])
            celda.border = borde
            if cc['tipo'] == 'participacion':
                celda.number_format = '0.0%'
            elif cc['tipo'] in FORMATOS:
                celda.number_format = FORMATOS[cc['tipo']]
            if cc['clave'] == 'remision' and datos.get('_id'):
                celda.font = Font(color=AZUL, bold=True)

    # Totales
    for i, cc in enumerate(columnas, start=1):
        clave, tipo = cc['clave'], cc['tipo']
        rango = f'{L(i)}{primera}:{L(i)}{ultima}'
        valor = None
        if i == 1:
            valor = (f'="Total: "&SUBTOTAL(3,{rango})&" remisiones"' if clave == 'remision' else 'Total')
        elif clave in ('merma_pct', 'diferencia_kg', 'promedio_carga_kg', 'promedio_descarga_kg') or clave.endswith('_ton'):
            valor = formula(clave, total)
        elif clave == '_participacion':
            valor = f'=SUBTOTAL(9,{rango})'
        if valor is None and tipo in ('entero', 'kg', 'ton', 'dinero', 'horas', 'pct'):
            valor = (f'=IFERROR(SUBTOTAL(1,{rango}),"")' if clave in PROMEDIADAS
                     else None if clave in NO_ADITIVAS else f'=SUBTOTAL(9,{rango})')
        celda = ws.cell(row=total, column=i, value=valor)
        celda.font = Font(bold=True)
        celda.fill = gris
        celda.number_format = '0.0%' if tipo == 'participacion' else FORMATOS.get(tipo, 'General')

    # Merma alta en rojo
    if 'merma_pct' in col and filas:
        letra = L(col['merma_pct'])
        ws.conditional_formatting.add(
            f'{letra}{primera}:{letra}{ultima}',
            CellIsRule(operator='greaterThan', formula=[str(UMBRAL_MERMA)],
                       font=Font(color=ROJO, bold=True),
                       fill=PatternFill('solid', start_color='FDECEA', end_color='FDECEA')))

    for i, cc in enumerate(columnas, start=1):
        largo = max([len(str(cc['titulo']))] + [len(str(d.get(cc['clave']) or '')) for d in filas[:300]])
        ws.column_dimensions[L(i)].width = max(11, min(48, largo + 3))
    ws.freeze_panes = ws.cell(row=primera, column=2 if agrupado else 1)
    if filas:
        ws.auto_filter.ref = f'A{enc}:{L(len(columnas))}{ultima}'

    # Gráfica de la tabla agrupada (una dimensión): serie de tiempo completa
    # (hasta 60 puntos) o los 20 grupos principales (la tabla ya viene ordenada)
    tiempo = agrupado and c['agrupar_por'][0] in R.DIMENSIONES_TIEMPO
    if agrupado and principal and len(c['agrupar_por']) == 1 and len(filas) > 1 and (not tiempo or len(filas) <= 60):
        n_graf = len(filas) if tiempo else min(len(filas), 20)
        grafica = LineChart() if tiempo else BarChart()
        if not tiempo:
            grafica.type = 'bar' if n_graf > 12 else 'col'
        prefijo = f'Top {n_graf}: ' if n_graf < len(filas) else ''
        grafica.title = f'{prefijo}{R.METRICAS[principal][0]} por {R.DIMENSIONES[c["agrupar_por"][0]].lower()}'
        grafica.legend = None
        grafica.height = max(8, min(22, 0.55 * n_graf + 4)) if not tiempo and n_graf > 12 else 9
        grafica.width = 20
        grafica.add_data(Reference(ws, min_col=col[principal], min_row=enc, max_row=enc + n_graf), titles_from_data=True)
        grafica.set_categories(Reference(ws, min_col=1, min_row=primera, max_row=enc + n_graf))
        if not tiempo and grafica.type == 'bar':
            grafica.y_axis.scaling.orientation = 'minMax'
            grafica.x_axis.scaling.orientation = 'maxMin'  # el mayor arriba
        ws.add_chart(grafica, f'{L(len(columnas) + 2)}{enc}')


def _hoja_totales(ws, res, enc):
    """Consulta de solo totales: tabla Indicador / Valor con fórmulas derivadas."""
    from openpyxl.styles import Font
    r = res['resumen']
    for i, titulo in enumerate(('Indicador', 'Valor'), start=1):
        _estilo_encabezado(ws.cell(row=enc, column=i, value=titulo))
    filas = [('Remisiones', r['remisiones'], 'entero', 'n'),
             ('Carga (kg)', r['carga_kg'], 'kg', 'c'),
             ('Descarga (kg)', r['descarga_kg'], 'kg', 'd'),
             ('Rechazo (kg)', r['rechazo_kg'], 'kg', 'r')]
    pos = {}
    for i, (etiqueta, valor, tipo, alias) in enumerate(filas, start=enc + 1):
        ws.cell(row=i, column=1, value=etiqueta)
        ws.cell(row=i, column=2, value=valor).number_format = FORMATOS[tipo]
        pos[alias] = f'B{i}'
    fila = enc + len(filas) + 1
    derivadas = [
        ('Carga (t)', f'={pos["c"]}/1000', 'ton'),
        ('Descarga (t)', f'={pos["d"]}/1000', 'ton'),
        ('Diferencia (kg)', f'={pos["d"]}+{pos["r"]}-{pos["c"]}', 'kg'),
        ('Merma %', f'=IF({pos["c"]}>0,({pos["c"]}-{pos["d"]}-{pos["r"]})/{pos["c"]}*100,0)', 'pct'),
        ('Carga promedio por remisión (kg)', f'=IF({pos["n"]}>0,{pos["c"]}/{pos["n"]},0)', 'kg'),
        ('Descarga promedio por remisión (kg)', f'=IF({pos["n"]}>0,{pos["d"]}/{pos["n"]},0)', 'kg'),
    ]
    for m in ('bultos', 'importe', 'peso_bascula_kg', 'tiempo_carga_h', 'tiempo_descarga_h', 'tiempo_transito_h'):
        if r.get(m):
            derivadas.append((R.METRICAS[m][0], r[m], R.METRICAS[m][1]))
    for etiqueta, valor, tipo in derivadas:
        ws.cell(row=fila, column=1, value=etiqueta)
        celda = ws.cell(row=fila, column=2, value=valor)
        celda.number_format = FORMATOS[tipo]
        if isinstance(valor, str) and valor.startswith('='):
            celda.font = Font(bold=True)
        fila += 1
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 22


def _hoja_resumen(ws, hojas, destinos, user, titulo, pregunta, respuesta):
    from openpyxl.styles import Alignment, Font, PatternFill

    _encabezado_hoja(ws, titulo or (hojas[0][1]['titulo'] if hojas else 'Reporte de remisiones'), user, [])
    fila = 4
    if pregunta:
        ws.cell(row=fila, column=1, value=f'Pregunta: {pregunta}').font = Font(italic=True, color='444444')
        fila += 1
    if respuesta:
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        celda = ws.cell(row=fila, column=1, value=_texto_plano(respuesta)[:3000])
        celda.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[fila].height = min(300, 15 * (1 + len(celda.value) // 95 + celda.value.count('\n')))
        fila += 1
    fila += 1

    claro = PatternFill('solid', start_color=GRIS, end_color=GRIS)
    for (c, res), destino in zip(hojas, destinos):
        r = res['resumen']
        celda = ws.cell(row=fila, column=1, value=res['titulo'])
        celda.font = Font(bold=True, size=12, color=AZUL, underline='single')
        celda.hyperlink = f"#'{destino}'!A1"
        fila += 1
        ws.cell(row=fila, column=1, value=' · '.join(res['filtros_texto'])).font = Font(size=9, color='666666')
        fila += 1
        for i, t in enumerate(('Indicador', 'Valor'), start=1):
            _estilo_encabezado(ws.cell(row=fila, column=i, value=t))
        fila += 1
        pos = {}
        base = [('Remisiones', r['remisiones'], 'entero', 'n'), ('Carga (kg)', r['carga_kg'], 'kg', 'c'),
                ('Descarga (kg)', r['descarga_kg'], 'kg', 'd'), ('Rechazo (kg)', r['rechazo_kg'], 'kg', 'r')]
        for etiqueta, valor, tipo, alias in base:
            ws.cell(row=fila, column=1, value=etiqueta)
            ws.cell(row=fila, column=2, value=valor).number_format = FORMATOS[tipo]
            pos[alias] = f'B{fila}'
            fila += 1
        for etiqueta, formula, tipo in (
                ('Carga (t)', f'={pos["c"]}/1000', 'ton'),
                ('Descarga (t)', f'={pos["d"]}/1000', 'ton'),
                ('Merma %', f'=IF({pos["c"]}>0,({pos["c"]}-{pos["d"]}-{pos["r"]})/{pos["c"]}*100,0)', 'pct')):
            ws.cell(row=fila, column=1, value=etiqueta).fill = claro
            celda = ws.cell(row=fila, column=2, value=formula)
            celda.number_format = FORMATOS[tipo]
            celda.font = Font(bold=True)
            celda.fill = claro
            fila += 1
        if r.get('importe'):
            ws.cell(row=fila, column=1, value='Importe')
            ws.cell(row=fila, column=2, value=r['importe']).number_format = FORMATOS['dinero']
            fila += 1
        fila += 1
    ws.column_dimensions['A'].width = 46
    ws.column_dimensions['B'].width = 22
    for letra in 'CDEF':
        ws.column_dimensions[letra].width = 16


# ════════════════════════════════════════════════════════════════════════════
# PRESENTACIÓN (PowerPoint)
# ════════════════════════════════════════════════════════════════════════════

def _fmt(valor, tipo):
    if valor is None or valor == '':
        return '—'
    if tipo == 'entero':
        return f'{valor:,.0f}'
    if tipo == 'kg':
        return f'{valor:,.2f}'
    if tipo == 'ton':
        return f'{valor:,.3f}'
    if tipo == 'pct':
        return f'{valor:,.2f} %'
    if tipo == 'horas':
        return f'{valor:,.2f} h'
    if tipo == 'dinero':
        return f'${valor:,.2f}'
    if tipo == 'sino':
        return 'Sí' if valor else 'No'
    if tipo == 'estatus':
        return R.STATUS.get(valor, str(valor))
    if tipo == 'fecha' and hasattr(valor, 'strftime'):
        return valor.strftime('%d/%m/%Y')
    if tipo == 'fechahora' and hasattr(valor, 'strftime'):
        return timezone.localtime(valor).strftime('%d/%m/%Y %H:%M')
    if tipo == 'hora' and hasattr(valor, 'strftime'):
        return valor.strftime('%H:%M')
    return str(valor)


def construir_presentacion(hojas, user, titulo='', pregunta='', respuesta='') -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    vacia = prs.slide_layouts[6]
    titulo = titulo or (hojas[0][1]['titulo'] if hojas else 'Reporte de remisiones')
    total = 2 + len(hojas)

    _portada(prs.slides.add_slide(vacia), titulo, pregunta, user)
    _diapositiva_respuesta(prs.slides.add_slide(vacia), respuesta, hojas, 2, total)
    for n, (c, res) in enumerate(hojas, start=3):
        _diapositiva_consulta(prs.slides.add_slide(vacia), c, res, n, total)

    salida = BytesIO()
    prs.save(salida)
    return salida.getvalue()


def _rgb(hexa):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexa)


def _rect(slide, x, y, w, h, color, linea=None):
    from pptx.enum.shapes import MSO_SHAPE
    forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    forma.fill.solid()
    forma.fill.fore_color.rgb = _rgb(color)
    if linea:
        forma.line.color.rgb = _rgb(linea)
    else:
        forma.line.fill.background()
    forma.shadow.inherit = False
    return forma


def _texto(slide, x, y, w, h, texto, size=14, bold=False, color='1F2937', italic=False, alinear=None):
    from pptx.enum.text import MSO_ANCHOR
    caja = slide.shapes.add_textbox(x, y, w, h)
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    if alinear is not None:
        p.alignment = alinear
    _runs(p, texto, size, bold, color, italic)
    return caja


def _runs(parrafo, texto, size, bold=False, color='1F2937', italic=False):
    """Agrega el texto al párrafo respetando **negritas**."""
    from pptx.util import Pt
    for parte in re.split(r'(\*\*[^*]+\*\*)', texto or ''):
        if not parte:
            continue
        negrita = parte.startswith('**') and parte.endswith('**')
        run = parrafo.add_run()
        run.text = parte[2:-2] if negrita else parte
        run.font.size = Pt(size)
        run.font.bold = bold or negrita
        run.font.italic = italic
        run.font.color.rgb = _rgb(AZUL_OSCURO if negrita and not bold else color)


def _encabezado(slide, titulo, subtitulo=''):
    from pptx.util import Inches
    _rect(slide, 0, 0, Inches(13.333), Inches(1.05), AZUL_OSCURO)
    _rect(slide, 0, Inches(1.05), Inches(13.333), Inches(0.06), AZUL)
    _texto(slide, Inches(0.5), Inches(0.18), Inches(12.3), Inches(0.6), titulo, size=26, bold=True, color='FFFFFF')
    if subtitulo:
        _texto(slide, Inches(0.5), Inches(1.22), Inches(12.3), Inches(0.4), subtitulo, size=11, color='6B7280')


def _pie(slide, numero, total):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches
    _texto(slide, Inches(0.5), Inches(7.0), Inches(9), Inches(0.3),
           'Fuente: remisiones del sistema 3R Recycling · fechas por inicio de descarga', size=9, color='9CA3AF')
    _texto(slide, Inches(11.3), Inches(7.0), Inches(1.5), Inches(0.3), f'{numero} / {total}',
           size=9, color='9CA3AF', alinear=PP_ALIGN.RIGHT)


def _kpis(slide, resumen, y, alto=None):
    from pptx.util import Inches
    if not resumen:
        return
    tarjetas = [('Remisiones', _fmt(resumen['remisiones'], 'entero'), ''),
                ('Carga', f'{resumen["carga_ton"]:,.2f} t', f'{resumen["carga_kg"]:,.2f} kg'),
                ('Descarga', f'{resumen["descarga_ton"]:,.2f} t', f'{resumen["descarga_kg"]:,.2f} kg'),
                ('Merma', f'{resumen["merma_pct"]:,.2f} %', f'Dif. {resumen["diferencia_kg"]:,.2f} kg')]
    if resumen.get('importe'):
        tarjetas.append(('Importe', f'${resumen["importe"]:,.2f}', ''))
    ancho_total, sep = 12.33, 0.2
    ancho = (ancho_total - sep * (len(tarjetas) - 1)) / len(tarjetas)
    alto = alto or Inches(1.15)
    for i, (etiqueta, valor, detalle) in enumerate(tarjetas):
        x = Inches(0.5 + i * (ancho + sep))
        _rect(slide, x, y, Inches(ancho), alto, 'F8FAFC', linea='E2E8F0')
        _texto(slide, x + Inches(0.18), y + Inches(0.1), Inches(ancho - 0.3), Inches(0.3),
               etiqueta.upper(), size=10, bold=True, color='64748B')
        alerta = etiqueta == 'Merma' and resumen['merma_pct'] > UMBRAL_MERMA
        _texto(slide, x + Inches(0.18), y + Inches(0.36), Inches(ancho - 0.3), Inches(0.5),
               valor, size=22, bold=True, color=ROJO if alerta else '0F172A')
        if detalle:
            _texto(slide, x + Inches(0.18), y + Inches(0.8), Inches(ancho - 0.3), Inches(0.3),
                   detalle, size=10, color='64748B')


def _portada(slide, titulo, pregunta, user):
    from pptx.util import Inches
    _rect(slide, 0, 0, Inches(13.333), Inches(7.5), AZUL_OSCURO)
    _rect(slide, Inches(0.8), Inches(2.35), Inches(1.2), Inches(0.08), AZUL)
    _texto(slide, Inches(0.8), Inches(1.6), Inches(11), Inches(0.5), '3R RECYCLING · REPORTES CON IA',
           size=14, bold=True, color='93C5FD')
    _texto(slide, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.8), titulo, size=40, bold=True, color='FFFFFF')
    if pregunta:
        _texto(slide, Inches(0.8), Inches(4.5), Inches(11.7), Inches(1), f'“{pregunta}”', size=18,
               italic=True, color='CBD5E1')
    nombre = user.get_full_name() or user.username
    _texto(slide, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.5),
           f'{nombre} · {timezone.localtime().strftime("%d/%m/%Y %H:%M")}', size=12, color='94A3B8')


def _diapositiva_respuesta(slide, respuesta, hojas, numero, total):
    from pptx.util import Inches, Pt
    _encabezado(slide, 'Respuesta')
    resumen = hojas[0][1]['resumen'] if hojas else None
    y_texto = Inches(1.45)
    if resumen and resumen['remisiones']:
        _kpis(slide, resumen, Inches(1.4))
        y_texto = Inches(2.85)
    caja = slide.shapes.add_textbox(Inches(0.5), y_texto, Inches(12.3), Inches(6.8) - y_texto)
    tf = caja.text_frame
    tf.word_wrap = True
    lineas = [l.strip() for l in (respuesta or '').split('\n') if l.strip()] or ['(sin texto)']
    tamano = 18 if len(respuesta or '') < 500 else 15 if len(respuesta or '') < 1100 else 12
    for i, linea in enumerate(lineas[:18]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        vineta = re.match(r'^([•*-]|\d+[.)])\s+', linea)
        if vineta:
            linea = '•  ' + linea[vineta.end():]
        _runs(p, re.sub(r'^#+\s*', '', linea), tamano, color='1F2937')
    _pie(slide, numero, total)


def _diapositiva_consulta(slide, c, res, numero, total):
    from pptx.util import Inches
    _encabezado(slide, res['titulo'], ' · '.join(res['filtros_texto']))
    tabla = res.get('tabla')
    resumen = res['resumen']
    if c['tipo'] == 'total' or not tabla or not tabla['filas']:
        _kpis(slide, resumen, Inches(2.0), alto=Inches(1.5))
        extras = [(R.METRICAS[m][0], _fmt(resumen.get(m), R.METRICAS[m][1]))
                  for m in ('bultos', 'promedio_carga_kg', 'promedio_descarga_kg', 'peso_bascula_kg',
                            'tiempo_carga_h', 'tiempo_descarga_h', 'tiempo_transito_h')
                  if resumen.get(m)]
        if extras:
            _tabla_pptx(slide, ['Indicador', 'Valor'], [list(e) for e in extras],
                        Inches(0.5), Inches(3.9), Inches(7), [Inches(4.8), Inches(2.2)])
        if not resumen['remisiones']:
            _texto(slide, Inches(0.5), Inches(3.8), Inches(12), Inches(0.6),
                   'No hay remisiones con estos filtros.', size=18, color='64748B')
    elif c['tipo'] == 'agrupado':
        _kpis(slide, resumen, Inches(1.7), alto=Inches(1.0))
        _agrupado_pptx(slide, c, tabla)
    else:
        _kpis(slide, resumen, Inches(1.7), alto=Inches(1.0))
        _lista_pptx(slide, tabla)
    _pie(slide, numero, total)


def _agrupado_pptx(slide, c, tabla):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Inches, Pt

    columnas = tabla['columnas']
    dims = [col for col in columnas if col['clave'].startswith('g_')]
    principal = next((m for m in PRINCIPALES if any(col['clave'] == m for col in columnas)), 'remisiones')
    tipo_p = R.METRICAS[principal][1]
    filas = tabla['filas']
    tiempo = c['agrupar_por'][0] in R.DIMENSIONES_TIEMPO
    tope = 24 if tiempo else 12
    visibles = filas[:tope]

    def etiqueta(f):
        return ' / '.join(str(f[d['clave']]) for d in dims)[:40]

    datos = CategoryChartData()
    datos.categories = [etiqueta(f) for f in visibles]
    datos.add_series(R.METRICAS[principal][0], [float(f[principal] or 0) for f in visibles],
                     number_format=FORMATOS.get(tipo_p, 'General').replace('"%"', '').replace('"$"', '$'))
    tipo_grafica = (XL_CHART_TYPE.LINE_MARKERS if tiempo
                    else XL_CHART_TYPE.BAR_CLUSTERED if len(visibles) > 6 else XL_CHART_TYPE.COLUMN_CLUSTERED)
    marco = slide.shapes.add_chart(tipo_grafica, Inches(0.5), Inches(2.85), Inches(7.6), Inches(4.05), datos)
    grafica = marco.chart
    grafica.has_legend = False
    grafica.font.size = Pt(10)
    serie = grafica.plots[0].series[0]
    if tipo_grafica != XL_CHART_TYPE.LINE_MARKERS:
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = _rgb(AZUL)
        grafica.plots[0].gap_width = 60
    else:
        serie.format.line.color.rgb = _rgb(AZUL)
    if len(visibles) <= 12:
        plot = grafica.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(9)
        plot.data_labels.number_format = '#,##0.0' if tipo_p in ('kg', 'ton', 'horas', 'pct') else '#,##0'
        plot.data_labels.number_format_is_linked = False
        if tipo_grafica != XL_CHART_TYPE.LINE_MARKERS:
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    if tipo_grafica == XL_CHART_TYPE.BAR_CLUSTERED:
        grafica.category_axis.reverse_order = True  # el mayor arriba

    total_p = sum(float(f[principal] or 0) for f in filas) or 0
    encabezados = [dims[0]['titulo'] if len(dims) == 1 else 'Grupo', R.METRICAS[principal][0], '% del total']
    renglones = [[etiqueta(f), _fmt(f[principal], tipo_p),
                  f'{(float(f[principal] or 0) / total_p * 100):,.1f} %' if total_p else '—'] for f in filas[:10]]
    _tabla_pptx(slide, encabezados, renglones, Inches(8.35), Inches(2.85), Inches(4.5),
                [Inches(2.2), Inches(1.4), Inches(0.9)])
    if len(filas) > tope:
        _texto(slide, Inches(0.5), Inches(6.7), Inches(7.6), Inches(0.3),
               f'Se muestran {tope} de {len(filas)} grupos; el Excel los trae todos.', size=9, color='64748B')


def _lista_pptx(slide, tabla):
    from pptx.util import Inches
    preferidas = ['remision', 'fecha_descarga', 'empresa', 'origen', 'destino', 'materiales',
                  'operador', 'descarga_kg', 'descarga_ton', 'merma_pct', 'status']
    disponibles = {col['clave']: col for col in tabla['columnas']}
    elegidas = [disponibles[k] for k in preferidas if k in disponibles][:7]
    for col in tabla['columnas']:  # lo que pidió el usuario y no está entre las preferidas
        if len(elegidas) >= 7:
            break
        if col not in elegidas:
            elegidas.append(col)
    filas = tabla['filas'][:12]
    renglones = [[_fmt(f.get(col['clave']), col['tipo'])[:34] for col in elegidas] for f in filas]
    ancho = 12.33 / len(elegidas)
    _tabla_pptx(slide, [col['titulo'] for col in elegidas], renglones, Inches(0.5), Inches(2.85),
                Inches(12.33), [Inches(ancho)] * len(elegidas), tamano=10)
    if tabla['total_filas'] > len(filas):
        _texto(slide, Inches(0.5), Inches(6.72), Inches(12), Inches(0.3),
               f'Se muestran {len(filas)} de {tabla["total_filas"]:,} remisiones; el Excel las trae todas.',
               size=9, color='64748B')


def _tabla_pptx(slide, encabezados, renglones, x, y, ancho, anchos, tamano=11):
    from pptx.util import Inches, Pt
    filas = len(renglones) + 1
    marco = slide.shapes.add_table(filas, len(encabezados), x, y, ancho, Inches(0.32) * filas)
    tabla = marco.table
    for i, w in enumerate(anchos):
        tabla.columns[i].width = w
    for j, texto in enumerate(encabezados):
        celda = tabla.cell(0, j)
        celda.fill.solid()
        celda.fill.fore_color.rgb = _rgb(AZUL_OSCURO)
        celda.text = ''
        _runs(celda.text_frame.paragraphs[0], str(texto), tamano, bold=True, color='FFFFFF')
    for i, renglon in enumerate(renglones, start=1):
        for j, texto in enumerate(renglon):
            celda = tabla.cell(i, j)
            celda.fill.solid()
            celda.fill.fore_color.rgb = _rgb('FFFFFF' if i % 2 else 'F1F5F9')
            celda.text = ''
            _runs(celda.text_frame.paragraphs[0], str(texto), tamano, color='1F2937')
            celda.margin_top = celda.margin_bottom = Pt(2)
    return tabla
