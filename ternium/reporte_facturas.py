"""
Reporte detallado de facturas (Asignación de factura › Facturas › "Reporte detallado").

Excel con dos hojas (abre en el Resumen):
- Resumen: una fila por factura con su desglose, para filtrar y sumar.
- Detalle: por cada factura, una fila por concepto del XML (CFDI) con la razón social,
  RFC y UUID, las remisiones aplicadas a la factura, la fecha de emisión, el inicio de
  carga de la remisión del folio que trae el concepto (si está en el sistema), la clave SAT con su descripción (catálogo del SAT de facturación), cantidad,
  unidad, valor unitario, importe e impuestos (IVA, IEPS y retenciones; solo salen las
  columnas que apliquen). Es una tabla de Excel con filtros y un renglón de totales que
  respeta lo filtrado; el desglose de cada factura está en el Resumen.

Los conceptos se leen otra vez del XML guardado, no de lo capturado. Las facturas sin
XML salen en un renglón con sus remisiones aplicadas y el desglose capturado a mano.
"""
import datetime as dt
import io
import re

from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.views.decorators.http import require_http_methods

from .asignacion_factura import (_api, _base, _facturas, _norm, _numero_folio, catalogo_sat as _catalogo,
                                 claves_de_folio, filtrar_facturas, remisiones_del_folio)
from .cfdi_xml import ErrorCFDI, leer_cfdi

MAX_FACTURAS = 1000
USO_CFDI = {'G01': 'Adquisición de mercancías', 'G03': 'Gastos en general', 'S01': 'Sin efectos fiscales',
            'I08': 'Otra maquinaria y equipo', 'CP01': 'Pagos'}
METODO_PAGO = {'PUE': 'Una sola exhibición', 'PPD': 'Parcialidades o diferido'}
FORMA_PAGO = {'01': 'Efectivo', '02': 'Cheque', '03': 'Transferencia', '04': 'Tarjeta de crédito',
              '28': 'Tarjeta de débito', '99': 'Por definir'}


def _con_nombre(catalogo, clave):
    return f'{clave} · {catalogo[clave]}' if clave in catalogo else (clave or '')


def _local(valor):
    return timezone.localtime(valor).replace(tzinfo=None) if valor else None


def datos_factura(f, user):
    """La factura con sus conceptos (del XML guardado) y la remisión de cada concepto."""
    cfdi, avisos = None, []
    if f.xml:
        try:
            with f.xml.open('rb') as archivo:
                cfdi = leer_cfdi(archivo.read())
        except (OSError, ValueError, ErrorCFDI):
            avisos.append('No se pudo leer el XML guardado: se usan los datos registrados.')
    conceptos = cfdi['conceptos'] if cfdi else list(f.conceptos or [])
    emision = parse_datetime(cfdi['fecha_emision']) if cfdi and cfdi.get('fecha_emision') else None
    if emision and timezone.is_aware(emision):
        emision = _local(emision)

    ligadas = list(f.remisiones.exclude(status='CANCELADO'))
    indice = {}
    for r in ligadas:
        for k in claves_de_folio(r):
            indice.setdefault(k, r)
    otras = None  # remisiones de la unidad de negocio, para folios del XML que no están ligados
    filas = []
    for c in conceptos:
        remisiones, notas = [], []
        for folio in c.get('folios') or []:
            r = indice.get(_norm(folio))
            if r is None:
                if otras is None:
                    otras = _base(user).filter(empresa_id=f.empresa_id)
                encontradas = remisiones_del_folio(otras, folio)
                r = encontradas[0] if encontradas else None
                if r is not None:
                    notas.append(f'El folio {folio} no está ligado a esta factura')
                else:
                    de_otra = remisiones_del_folio(_base(user).exclude(empresa_id=f.empresa_id), folio)
                    notas.append(f'El folio {folio} es de {de_otra[0].empresa.nombre}, no de {f.empresa.nombre}' if de_otra
                                 else f'El folio {folio} no existe en el sistema')
            if r is not None and all(x.pk != r.pk for x in remisiones):
                remisiones.append(r)
        if not c.get('folios'):
            notas.append('La descripción no trae folio de remisión')
        filas.append({'concepto': c, 'inicio': _local(remisiones[0].inicia_ld) if remisiones else None,
                      'nota': '; '.join(notas)})
    if not conceptos:  # capturada sin XML: un renglón con sus remisiones aplicadas
        filas.append({'concepto': None, 'inicio': None, 'nota': ''})
    aplicadas = sorted(ligadas, key=lambda r: (_numero_folio(r.remision) or 0, r.remision))

    suma_conceptos = round(sum(c.get('importe', 0) for c in conceptos), 2)
    if conceptos and abs(suma_conceptos - float(f.subtotal)) > 0.01:
        avisos.append(f'Los conceptos suman ${suma_conceptos:,.2f} y el subtotal es ${float(f.subtotal):,.2f}.')
    return {
        'factura': f, 'cfdi': cfdi, 'emision': emision or f.fecha, 'filas': filas, 'avisos': avisos,
        'num_conceptos': len(conceptos),
        'aplicadas': ', '.join(r.remision for r in aplicadas),
        'receptor_nombre': (cfdi['receptor']['nombre'] if cfdi else '') or f.cliente,
        'receptor_rfc': (cfdi['receptor']['rfc'] if cfdi else '') or f.receptor_rfc,
    }


# ── Columnas de la hoja Detalle (las marcadas con "opcional" solo salen si algún renglón las usa) ──
COLUMNAS = [
    # clave, título, ancho, formato, opcional
    ('factura', 'Factura', 12, 'texto', None),
    ('razon_social', 'Razón social', 30, 'largo', None),
    ('rfc', 'RFC', 15, 'texto', None),
    ('uuid', 'UUID', 38, 'texto', None),
    ('aplicadas', 'Remisiones aplicadas', 26, 'largo', None),
    ('emision', 'Emisión de la factura', 17, 'fecha', None),
    ('inicio', 'Inicio de carga', 17, 'fecha', None),
    ('descripcion', 'Descripción (factura)', 46, 'largo', None),
    ('clave_sat', 'Clave SAT', 11, 'texto', None),
    ('clave_sat_desc', 'Clave SAT (descripción)', 24, 'largo', None),
    ('no_ident', 'No. identificación', 14, 'texto', None),
    ('cantidad', 'Cantidad', 11, 'cantidad', None),
    ('clave_unidad', 'Clave unidad', 9, 'texto', None),
    ('unidad', 'Unidad', 12, 'texto', None),
    ('valor_unitario', 'Valor unitario', 12, 'unitario', None),
    ('importe', 'Importe', 14, 'dinero', None),
    ('descuento', 'Descuento', 13, 'dinero', 'descuento'),
    ('iva_tasa', 'IVA %', 8, 'tasa', None),
    ('iva', 'IVA', 13, 'dinero', None),
    ('ieps_tasa', 'IEPS %', 8, 'tasa', 'ieps'),
    ('ieps', 'IEPS', 13, 'dinero', 'ieps'),
    ('ret_iva_tasa', 'IVA retenido %', 9, 'tasa', 'ret_iva'),
    ('ret_iva', 'IVA retenido', 13, 'dinero', 'ret_iva'),
    ('ret_isr_tasa', 'ISR retenido %', 9, 'tasa', 'ret_isr'),
    ('ret_isr', 'ISR retenido', 13, 'dinero', 'ret_isr'),
    ('total', 'Total', 15, 'dinero', None),
    ('moneda', 'Moneda', 9, 'texto', 'moneda'),
    ('nota', 'Observaciones', 30, 'largo', None),
]
DINERO = '"$"#,##0.00;[Red]-"$"#,##0.00'


def _aplican(datos):
    """Qué columnas opcionales hacen falta: impuestos que aparecen y moneda si hay pesos y dólares."""
    hay = {'moneda'} if len({d['factura'].moneda for d in datos}) > 1 else set()
    for d in datos:
        f = d['factura']
        for clave, campo in (('descuento', 'descuento'), ('ieps', 'ieps_importe'), ('ret_iva', 'ret_iva_importe'),
                             ('ret_isr', 'ret_isr_importe')):
            if getattr(f, campo):
                hay.add(clave)
        for fila in d['filas']:
            c = fila['concepto'] or {}
            for clave in ('descuento', 'ieps', 'ret_iva', 'ret_isr'):
                if c.get(clave):
                    hay.add(clave)
    return hay


def _valores(d, fila, cat_prod, cat_unidad):
    f, c = d['factura'], fila['concepto']
    v = {
        'factura': f.numero, 'razon_social': d['receptor_nombre'], 'rfc': d['receptor_rfc'], 'uuid': f.uuid,
        'emision': d['emision'], 'aplicadas': d['aplicadas'], 'inicio': fila['inicio'], 'nota': fila['nota'],
        'moneda': f.moneda,
    }
    if c is None:
        v['descripcion'] = '(Factura capturada sin XML)'
        return v
    importe, descuento = c.get('importe', 0), c.get('descuento', 0)
    total = c.get('total')
    if total is None:  # conceptos guardados antes del desglose por impuesto
        total = importe - descuento + c.get('traslados', 0) - c.get('retenciones', 0)
    v.update({
        'descripcion': c.get('descripcion', ''), 'clave_sat': c.get('clave_prod_serv', ''),
        'clave_sat_desc': cat_prod.get(c.get('clave_prod_serv', ''), ''), 'no_ident': c.get('no_identificacion', ''),
        'cantidad': c.get('cantidad'), 'clave_unidad': c.get('clave_unidad', ''),
        'unidad': c.get('unidad') or cat_unidad.get(c.get('clave_unidad', ''), ''),
        'valor_unitario': c.get('valor_unitario'), 'importe': importe, 'descuento': descuento,
        'iva_tasa': c.get('iva_tasa', 0), 'iva': c.get('iva', c.get('traslados', 0)),
        'ieps_tasa': c.get('ieps_tasa', 0), 'ieps': c.get('ieps', 0),
        'ret_iva_tasa': c.get('ret_iva_tasa', 0), 'ret_iva': c.get('ret_iva', 0),
        'ret_isr_tasa': c.get('ret_isr_tasa', 0), 'ret_isr': c.get('ret_isr', 0), 'total': total,
    })
    return v


def construir(datos, subtitulo):
    """Bytes del Excel (hojas Detalle y Resumen)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.table import Table, TableStyleInfo

    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    izquierda = Alignment(horizontal='left', vertical='center', wrap_text=True)
    derecha = Alignment(horizontal='right', vertical='center')
    azul, tinta = PatternFill('solid', fgColor='0B1730'), Font(bold=True, color='FFFFFF')

    def formato(celda, tipo, valor):
        if tipo == 'dinero':
            celda.number_format = DINERO
        elif tipo == 'unitario':
            celda.number_format = '"$"#,##0.00##'
        elif tipo == 'kg':
            celda.number_format = '#,##0'
        elif tipo == 'cantidad':
            celda.number_format = '#,##0' if valor is None or float(valor).is_integer() else '#,##0.000'
        elif tipo == 'tasa':
            celda.number_format = '0" %"' if valor is None or float(valor).is_integer() else '0.0000" %"'
        elif tipo == 'fecha':
            celda.number_format = 'dd/mm/yyyy hh:mm' if isinstance(valor, dt.datetime) else 'dd/mm/yyyy'
        celda.alignment = izquierda if tipo == 'largo' else centro

    cat_prod = _catalogo('ClaveProdServ', {(fl['concepto'] or {}).get('clave_prod_serv') for d in datos for fl in d['filas']})
    cat_unidad = _catalogo('ClaveUnidad', {(fl['concepto'] or {}).get('clave_unidad') for d in datos for fl in d['filas']})
    hay = _aplican(datos)
    columnas = [col for col in COLUMNAS if col[4] is None or col[4] in hay]
    n_cols = len(columnas)
    pos = {col[0]: i for i, col in enumerate(columnas, 1)}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Detalle'
    ws.sheet_view.showGridLines = False

    def titulo(hoja, texto, ancho):
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
        hoja.cell(row=1, column=1, value=texto).font = Font(bold=True, size=14, color='FFFFFF')
        hoja.cell(row=1, column=1).fill = azul
        hoja.cell(row=1, column=1).alignment = centro
        hoja.row_dimensions[1].height = 28
        hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
        hoja.cell(row=2, column=1, value=subtitulo).font = Font(italic=True, size=10, color='475569')
        hoja.cell(row=2, column=1).alignment = centro

    titulo(ws, 'Reporte detallado de facturas · 3R Recycling', n_cols)
    enc = 4
    for i, (_, texto, ancho, _, _) in enumerate(columnas, 1):
        celda = ws.cell(row=enc, column=i, value=texto)
        celda.font, celda.fill, celda.alignment = tinta, PatternFill('solid', fgColor='1D4ED8'), centro
        ws.column_dimensions[L(i)].width = ancho
    ws.row_dimensions[enc].height = 32

    # Una fila por concepto, sin renglones intercalados: es una tabla de Excel con filtros.
    # El desglose de cada factura está en el Resumen.
    fila_n = enc + 1
    for d in datos:
        for n, fila in enumerate(d['filas']):
            valores = _valores(d, fila, cat_prod, cat_unidad)
            if n == 0 and d['avisos']:  # avisos de la factura (p. ej. los conceptos no cuadran con el subtotal)
                valores['nota'] = '; '.join(d['avisos'] + ([valores['nota']] if valores.get('nota') else []))
            for clave, _, _, tipo, _ in columnas:
                valor = valores.get(clave)
                celda = ws.cell(row=fila_n, column=pos[clave], value=valor if valor not in ('',) else None)
                formato(celda, tipo, valor)
                if clave == 'nota' and valor:
                    celda.font = Font(color='B45309', italic=True)
            fila_n += 1
    ultima_d = fila_n - 1
    tabla_d = Table(displayName='DetalleFacturas', ref=f'A{enc}:{L(n_cols)}{max(ultima_d, enc + 1)}')
    tabla_d.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tabla_d)
    # Totales de lo que se ve (SUBTOTAL respeta los filtros de la tabla)
    importes = [pos[k] for k in ('importe', 'descuento', 'iva', 'ieps', 'ret_iva', 'ret_isr', 'total') if k in pos]
    fila_t = ultima_d + 1
    monedas = {d['factura'].moneda for d in datos}
    ws.cell(row=fila_t, column=importes[0] - 1,
            value='Total de lo filtrado' + ('' if len(monedas) == 1 else ' (filtra una moneda)')).font = Font(bold=True)
    ws.cell(row=fila_t, column=importes[0] - 1).alignment = derecha
    for col in importes:
        celda = ws.cell(row=fila_t, column=col, value=f'=SUBTOTAL(109,{L(col)}{enc + 1}:{L(col)}{ultima_d})')
        celda.number_format, celda.font, celda.alignment = DINERO, Font(bold=True), centro
        celda.border = Border(top=Side(style='medium', color='0B1730'))

    ws.freeze_panes = f'B{enc + 1}'
    ws.print_title_rows = f'{enc}:{enc}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ── Hoja Resumen: una fila por factura ──
    rs = wb.create_sheet('Resumen')
    rs.sheet_view.showGridLines = False
    resumen = [
        ('Captura', 11, 'texto'), ('Factura', 12, 'texto'), ('Emisión', 17, 'fecha'), ('Unidad de negocio', 18, 'texto'), ('Cliente', 30, 'largo'),
        ('RFC receptor', 15, 'texto'), ('UUID', 38, 'texto'), ('Uso CFDI', 22, 'largo'), ('Método de pago', 20, 'largo'),
        ('Forma de pago', 18, 'largo'), ('Moneda', 9, 'texto'), ('Remisiones aplicadas', 26, 'largo'), ('Conceptos', 10, 'kg'),
        ('Subtotal', 14, 'dinero'), ('Descuento', 13, 'dinero'), ('IEPS', 13, 'dinero'),
        ('IVA trasladado', 14, 'dinero'), ('IVA retenido', 14, 'dinero'), ('ISR retenido', 13, 'dinero'),
        ('Total', 15, 'dinero'), ('Con XML', 9, 'texto'),
    ]
    titulo(rs, 'Resumen de facturas · 3R Recycling', len(resumen))
    for i, (texto, ancho, _) in enumerate(resumen, 1):
        celda = rs.cell(row=enc, column=i, value=texto)
        celda.font, celda.fill, celda.alignment = tinta, PatternFill('solid', fgColor='1D4ED8'), centro
        rs.column_dimensions[L(i)].width = ancho
    for n, d in enumerate(datos, enc + 1):
        f = d['factura']
        valores = [f.numero_captura, f.numero, d['emision'], f.empresa.nombre, d['receptor_nombre'], d['receptor_rfc'], f.uuid,
                   _con_nombre(USO_CFDI, f.uso_cfdi), _con_nombre(METODO_PAGO, f.metodo_pago),
                   _con_nombre(FORMA_PAGO, f.forma_pago), f.moneda, d['aplicadas'], d['num_conceptos'],
                   float(f.subtotal), float(f.descuento), float(f.ieps_importe), float(f.iva_importe),
                   float(f.ret_iva_importe), float(f.ret_isr_importe), float(f.total), 'Sí' if f.xml else 'No']
        for i, ((_, _, tipo), valor) in enumerate(zip(resumen, valores), 1):
            celda = rs.cell(row=n, column=i, value=valor if valor != '' else None)
            formato(celda, tipo, valor)
    ultima = enc + len(datos)
    tabla = Table(displayName='ResumenFacturas', ref=f'A{enc}:{L(len(resumen))}{ultima}')
    tabla.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    rs.add_table(tabla)
    # Totales por moneda (no se suman pesos con dólares)
    titulos = [t for t, _, _ in resumen]
    col_moneda = L(titulos.index('Moneda') + 1)
    dinero = [i for i, (_, _, tipo) in enumerate(resumen, 1) if tipo == 'dinero']
    fila_t = ultima + 2
    for moneda in sorted({d['factura'].moneda for d in datos}):
        rs.cell(row=fila_t, column=dinero[0] - 1, value=f'Total {moneda}').font = Font(bold=True)
        rs.cell(row=fila_t, column=dinero[0] - 1).alignment = derecha
        for col in dinero:
            letra = L(col)
            celda = rs.cell(row=fila_t, column=col,
                            value=f'=SUMIFS({letra}{enc + 1}:{letra}{ultima},{col_moneda}{enc + 1}:{col_moneda}{ultima},"{moneda}")')
            celda.number_format, celda.font, celda.alignment = DINERO, Font(bold=True), centro
        fila_t += 1
    rs.freeze_panes = f'C{enc + 1}'  # captura y factura fijas
    # Que el resumen se vea completo: cada columna al ancho de su contenido
    for i, (texto, ancho, tipo) in enumerate(resumen, 1):
        largos = [len(str(v)) for v in (rs.cell(row=n, column=i).value for n in range(enc + 1, ultima + 1))
                  if v is not None and not isinstance(v, (dt.date, dt.datetime))]
        if tipo == 'dinero':
            largos = [len(f'${float(rs.cell(row=n, column=i).value or 0):,.2f}') for n in range(enc + 1, ultima + 1)]
        rs.column_dimensions[L(i)].width = min(max([ancho, len(texto) + 4] + [n + 3 for n in largos]), 60)
    rs.page_setup.orientation = 'landscape'
    rs.page_setup.fitToWidth, rs.page_setup.fitToHeight = 1, 0
    rs.sheet_properties.pageSetUpPr.fitToPage = True

    # El resumen es la primera pestaña (la que abre el archivo); después el detalle
    wb.move_sheet(rs, offset=-wb.index(rs))
    wb.active = 0
    rs.sheet_view.tabSelected, ws.sheet_view.tabSelected = True, False

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


@require_http_methods(['GET'])
@_api
def api_reporte(request):
    """GET ?q&empresa&cliente&desde&hasta (los filtros de la lista) o ?factura=<id> → Excel del reporte detallado."""
    p = request.GET
    qs = _facturas(request.user)
    if str(p.get('factura', '')).isdigit():
        qs = qs.filter(pk=int(p['factura']))
    else:
        qs = filtrar_facturas(qs, p)
    qs = qs.order_by('fecha', 'numero')
    cuantas = qs.count()
    if not cuantas:
        return JsonResponse({'detail': 'No hay facturas con esos filtros.'}, status=404)
    if cuantas > MAX_FACTURAS:
        return JsonResponse({'detail': f'Son {cuantas:,} facturas: acota por fechas o unidad de negocio (máximo {MAX_FACTURAS:,}).'},
                            status=400)
    datos = [datos_factura(f, request.user) for f in qs]

    filtros = []
    if p.get('desde') or p.get('hasta'):
        filtros.append(f'del {p.get("desde") or "inicio"} al {p.get("hasta") or "hoy"}')
    if p.get('q'):
        filtros.append(f'búsqueda "{p["q"]}"')
    if str(p.get('empresa', '')).isdigit():
        filtros.append(datos[0]['factura'].empresa.nombre)
    if (p.get('cliente') or '').strip():
        filtros.append(f'cliente {p["cliente"].strip()}')
    subtitulo = (f'{cuantas} factura(s)' + (' · ' + ' · '.join(filtros) if filtros else '')
                 + f' · generado el {timezone.localtime():%d/%m/%Y %H:%M}')
    contenido = construir(datos, subtitulo)
    if cuantas == 1:
        nombre = 'Reporte factura ' + re.sub(r'[^A-Za-z0-9_-]', '_', datos[0]['factura'].numero)
    else:
        nombre = f'Reporte facturas {timezone.localdate():%Y-%m-%d}'
    respuesta = HttpResponse(contenido, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}.xlsx"'
    return respuesta
