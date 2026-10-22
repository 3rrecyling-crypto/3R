"""
Excel de una tabla que ya se ve en pantalla (p. ej. Catálogos). Next manda los títulos de las
columnas y las filas tal como las tiene (con la búsqueda y el orden aplicados) y aquí se arma
el archivo con el estilo de los demás reportes: una tabla de Excel con filtros.

No consulta datos: solo da formato a lo que el usuario ya tiene en pantalla.
"""
import io
import json
import re

from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.utils.http import content_disposition_header
from django.views.decorators.http import require_http_methods

MAX_FILAS = 20000
MAX_COLUMNAS = 80
MAX_TEXTO = 32767  # lo más que cabe en una celda de Excel
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _valor(v):
    """Número, texto o vacío. Un texto nunca se vuelve fórmula aunque empiece con "="."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    if v is None or v == '':
        return None
    if isinstance(v, bool):
        return 'Sí' if v else 'No'
    if isinstance(v, (int, float)):
        return v
    return ILLEGAL_CHARACTERS_RE.sub('', str(v))[:MAX_TEXTO]


def construir(titulo, columnas, filas):
    """Bytes del .xlsx: una hoja con la tabla (encabezado fijo y filtros)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\[\]:*?/\\]', ' ', titulo)[:31] or 'Datos'
    ws.sheet_view.showGridLines = False

    # Los encabezados de una tabla de Excel no se pueden repetir
    vistos, encabezados = {}, []
    for c in columnas:
        texto = ' '.join(str(c).split()) or 'Columna'
        vistos[texto] = vistos.get(texto, 0) + 1
        encabezados.append(texto if vistos[texto] == 1 else f'{texto} ({vistos[texto]})')

    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, texto in enumerate(encabezados, 1):
        celda = ws.cell(row=1, column=i, value=texto)
        celda.font, celda.fill, celda.alignment = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1D4ED8'), centro
    ws.row_dimensions[1].height = 30

    n = len(encabezados)
    largos = [len(t) + 4 for t in encabezados]
    izquierda = Alignment(vertical='center', wrap_text=True)
    for r, fila in enumerate(filas, 2):
        for i in range(n):
            valor = _valor(fila[i] if i < len(fila) else None)
            celda = ws.cell(row=r, column=i + 1, value=valor)
            if isinstance(valor, str):
                celda.data_type = 's'  # texto tal cual (nada de fórmulas)
                celda.alignment = izquierda
                largos[i] = max(largos[i], min(len(valor), 60) + 2)
            elif valor is not None:
                largos[i] = max(largos[i], len(str(valor)) + 2)
    for i, ancho in enumerate(largos, 1):
        ws.column_dimensions[L(i)].width = min(max(ancho, 8), 60)

    tabla = Table(displayName='Datos', ref=f'A1:{L(n)}{max(len(filas) + 1, 2)}')
    tabla.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = 'A2'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


@require_http_methods(['POST'])
def api_excel_tabla(request):
    """POST JSON {titulo, columnas: [texto], filas: [[valor, ...]]} → .xlsx con esa tabla."""
    if not request.user.is_authenticated:
        return JsonResponse({'detail': 'No autenticado.'}, status=401)
    try:
        body = json.loads(request.body or b'{}')
    except ValueError:
        return JsonResponse({'detail': 'Datos inválidos.'}, status=400)
    titulo = ' '.join(str(body.get('titulo') or 'Tabla').split())[:60]
    columnas, filas = body.get('columnas'), body.get('filas')
    if not isinstance(columnas, list) or not columnas or len(columnas) > MAX_COLUMNAS:
        return JsonResponse({'detail': f'Faltan las columnas (máximo {MAX_COLUMNAS}).'}, status=400)
    if not isinstance(filas, list) or not all(isinstance(f, list) for f in filas):
        return JsonResponse({'detail': 'Las filas no vienen bien.'}, status=400)
    if len(filas) > MAX_FILAS:
        return JsonResponse({'detail': f'Son {len(filas):,} filas: el máximo es {MAX_FILAS:,}. Usa la búsqueda.'}, status=400)
    respuesta = HttpResponse(construir(titulo, columnas, filas), content_type=XLSX)
    respuesta['Content-Disposition'] = content_disposition_header(True, f'{titulo} {timezone.localdate():%Y-%m-%d}.xlsx')
    return respuesta
